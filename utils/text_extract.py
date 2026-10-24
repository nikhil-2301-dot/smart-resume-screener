"""
Stage 2 & 3: Upload multiple PDF/DOCX resumes and extract raw text.
"""
import io
import fitz  # PyMuPDF
import docx  # python-docx


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract all text from a PDF file given as raw bytes."""
    text_parts = []
    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        for page in doc:
            text_parts.append(page.get_text())
    return "\n".join(text_parts)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract all text from a DOCX file given as raw bytes."""
    file_stream = io.BytesIO(file_bytes)
    document = docx.Document(file_stream)

    parts = [p.text for p in document.paragraphs]

    # Also pull text out of tables, since resumes sometimes use them for layout
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)

    return "\n".join(parts)


def extract_text(filename: str, file_bytes: bytes) -> str:
    """Dispatch to the right extractor based on file extension."""
    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif lower_name.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: {filename}. Only PDF and DOCX are supported.")
